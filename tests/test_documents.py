import io
from datetime import date, datetime

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader

from turing_projects_v2.applications import ApplicationRow, ApplicationTable
from turing_projects_v2.documents import application_table_pdf, letter_docx, letter_pdf

LETTER = (
    "Ada Example\n"
    "Sample Town, 12345\n"
    "\n"
    "Dear Mr/Ms Hiring Manager,\n"
    "I apply for Sr. AI/ML Engineer – (Remote) at R&D <Labs> 李.\n"  # noqa: RUF001 (the en dash is on purpose)
    "Yours sincerely,\n"
    "Ada Example\n"
)


def test_docx_holds_the_letter_line_by_line():
    document = docx.Document(io.BytesIO(letter_docx(LETTER)))

    assert [paragraph.text for paragraph in document.paragraphs] == LETTER.splitlines()


def test_pdf_is_a_pdf_document_even_with_unusual_characters():
    pdf = letter_pdf(LETTER)

    assert pdf.startswith(b"%PDF-")
    assert b"/Type /Page" in pdf


DIN_LETTER = (
    "Ada Example\n"
    "Example Street 1\n"
    "12345 Sample Town\n"
    "Germany\n"
    "Tel.: 0123 4567890\n"
    "ada@example.com\n"
    "\n"
    "Acme Corp\n"
    "1 Main St, Austin, TX, US\n"
    "\n"
    "Sample Town, 25 October 2026\n"
    "\n"
    "Application for the position of AI Engineer\n"
    "\n"
    "Dear Mr/Ms Hiring Manager,\n"
    "I am writing to apply for AI Engineer at Acme Corp.\n"
    "My training covers the following topics:\n"
    "- AI agents\n"
    "- Evaluation benchmarks\n"
    "I look forward to hearing from you.\n"
    "Yours sincerely,\n"
    "Ada Example\n"
)


def paragraphs_by_text(document):
    return {paragraph.text: paragraph for paragraph in document.paragraphs if paragraph.text}


def test_docx_lays_the_letter_out_as_a_german_business_letter():
    document = docx.Document(io.BytesIO(letter_docx(DIN_LETTER)))
    paragraphs = paragraphs_by_text(document)
    texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]

    assert texts[:6] == DIN_LETTER.splitlines()[:6]  # the sender block comes first
    assert paragraphs["12345 Sample Town"].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraphs["Acme Corp"].alignment in (None, WD_ALIGN_PARAGRAPH.LEFT)
    assert paragraphs["Sample Town, 25 October 2026"].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    subject = paragraphs["Application for the position of AI Engineer"]
    assert all(run.bold for run in subject.runs)
    body = paragraphs["I am writing to apply for AI Engineer at Acme Corp."]
    assert body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert body.paragraph_format.line_spacing == 1.5
    bullet = paragraphs["AI agents"]
    assert bullet.style.name == "List Bullet"
    assert texts[-2:] == ["Yours sincerely,", "Ada Example"]


def pdf_text(pdf):
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)


def application_row(day, employer, **changes):
    row = {
        "applied_on": day,
        "employer": employer,
        "position": "Wissenschaftliche Mitarbeiterin (m/w/d)",
        "city": "Beispielstadt",
        "employment": "Vollzeit",
        "work_mode": "Präsenz",
    }
    return ApplicationRow(**(row | changes))


LONG_EMPLOYER = "Größere Forschungsgesellschaft für angewandte Künstliche Intelligenz mbH"


def test_application_table_pdf_is_a_german_table_of_the_applications_in_order():
    table = ApplicationTable(
        candidate_name="Ada Example",
        made_at=datetime(2026, 11, 30, 9, 15),
        rows=[
            application_row(date(2026, 11, 2), "Universität Beispielstadt"),
            application_row(
                date(2026, 11, 28),
                LONG_EMPLOYER,
                city="nicht angegeben",
                employment="Teilzeit",
                work_mode="nicht angegeben",
            ),
        ],
    )

    text = pdf_text(application_table_pdf(table))

    assert text.startswith("Bewerbungen\nAda Example\nStand: 30.11.2026")
    header = "Datum Arbeitgeber Stelle Ort Vollzeit/Teilzeit Präsenz/Online"
    assert header in " ".join(text.split())
    first, second = text.index("02.11.2026"), text.index("28.11.2026")
    assert text.index("Datum") < first < second
    assert "Universität Beispielstadt" in text[first:second]
    assert "Teilzeit" in text[second:]
    assert text[second:].count("nicht angegeben") == 2
    # A long employer wraps inside its cell instead of running off the page.
    assert not any(LONG_EMPLOYER in line for line in text.splitlines())
    assert "Größere" in text[second:] and "mbH" in text[second:]
