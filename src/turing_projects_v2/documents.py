"""Cover Letter downloads: the letter as a .docx and as a PDF, laid out as a German business
letter (DIN 5008): sender block on the right, recipient on the left, place and date on the right,
a bold subject line, then the body justified with 1.5 line spacing.

The layout is read from the letter text, whose blocks are separated by blank lines: sender,
recipient, place and date, subject, body. A letter without those blocks (saved before the
layout existed) is written line by line.
"""

import io
from xml.sax.saxutils import escape

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from pydantic import BaseModel
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

FONT_SIZE = 11
LINE_SPACING = 1.5
LIST_LINE_SPACING = 1.2  # bullet lists are set tighter, so a full letter fits one page
BULLET = "- "
# DIN 5008 page margins.
MARGINS = {"left": 2.5, "right": 2.0, "top": 1.7, "bottom": 1.5}  # cm
# Space (pt) before each block, small enough for a one-page letter.
GAPS = {
    "recipient": 20,
    "date": 16,
    "subject": 16,
    "salutation": 14,
    "paragraph": 6,
    "closing": 12,
    "signature": 26,
}


class LetterLayout(BaseModel):
    sender: list[str]
    recipient: list[str]
    date: str
    subject: str
    salutation: str
    body: list[str]
    """Paragraphs and "- " bullet lines between the salutation and the closing."""
    closing: str
    signature: str


def letter_layout(text: str) -> LetterLayout | None:
    """The letter's blocks, or None if the text does not have them."""
    blocks = [block.strip("\n").splitlines() for block in text.strip("\n").split("\n\n")]
    blocks = [block for block in blocks if block]
    if len(blocks) < 5 or len(blocks[2]) != 1 or len(blocks[3]) != 1:
        return None
    sender, recipient, (date,), (subject,), *rest = blocks
    body = [line for block in rest for line in block]
    if len(body) < 3:
        return None
    return LetterLayout(
        sender=sender,
        recipient=recipient,
        date=date,
        subject=subject,
        salutation=body[0],
        body=body[1:-2],
        closing=body[-2],
        signature=body[-1],
    )


def letter_docx(text: str) -> bytes:
    """The letter as a Word document."""
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(FONT_SIZE)
    style.paragraph_format.space_after = Pt(0)
    for section in document.sections:
        section.left_margin = Cm(MARGINS["left"])
        section.right_margin = Cm(MARGINS["right"])
        section.top_margin = Cm(MARGINS["top"])
        section.bottom_margin = Cm(MARGINS["bottom"])
    layout = letter_layout(text)
    if layout is None:
        for line in text.splitlines():
            document.add_paragraph(line)
    else:
        _docx_letter(document, layout)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_letter(document: docx.document.Document, layout: LetterLayout) -> None:
    def add(text: str, alignment=None, space_before: float = 0, bold: bool = False):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        run.bold = bold
        if alignment is not None:
            paragraph.alignment = alignment
        paragraph.paragraph_format.space_before = Pt(space_before)
        return paragraph

    for line in layout.sender:
        add(line, WD_ALIGN_PARAGRAPH.RIGHT)
    for number, line in enumerate(layout.recipient):
        add(line, space_before=GAPS["recipient"] if number == 0 else 0)
    add(layout.date, WD_ALIGN_PARAGRAPH.RIGHT, space_before=GAPS["date"])
    add(layout.subject, space_before=GAPS["subject"], bold=True)
    add(layout.salutation, space_before=GAPS["salutation"])
    for line in layout.body:
        if line.startswith(BULLET):
            paragraph = document.add_paragraph(line.removeprefix(BULLET), style="List Bullet")
            paragraph.paragraph_format.line_spacing = LIST_LINE_SPACING
        else:
            paragraph = document.add_paragraph(line)
            paragraph.paragraph_format.space_before = Pt(GAPS["paragraph"])
            paragraph.paragraph_format.line_spacing = LINE_SPACING
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add(layout.closing, space_before=GAPS["closing"])
    add(layout.signature, space_before=GAPS["signature"])


def letter_pdf(text: str) -> bytes:
    """The letter as an A4 PDF in Helvetica (Arial's metric twin, built into every PDF reader).

    Helvetica covers Western European text (umlauts, ß, dashes, curly quotes); other characters
    show as blanks.
    """
    layout = letter_layout(text)
    flowables = _pdf_lines(text) if layout is None else _pdf_letter(layout)
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=MARGINS["left"] * cm,
        rightMargin=MARGINS["right"] * cm,
        topMargin=MARGINS["top"] * cm,
        bottomMargin=MARGINS["bottom"] * cm,
        title="Cover Letter",
    )
    document.build(flowables)
    return buffer.getvalue()


def _style(name: str, alignment: int = TA_LEFT, **options) -> ParagraphStyle:
    return ParagraphStyle(
        name,
        fontName=options.pop("fontName", "Helvetica"),
        fontSize=FONT_SIZE,
        leading=options.pop("leading", FONT_SIZE * 1.2),
        alignment=alignment,
        **options,
    )


def _pdf_lines(text: str) -> list:
    style = _style("line", leading=14)
    # An empty paragraph has no height, so a blank line is kept as a non-breaking space.
    return [Paragraph(escape(line) or "&nbsp;", style) for line in text.splitlines()]


def _pdf_letter(layout: LetterLayout) -> list:
    right = _style("right", TA_RIGHT)
    left = _style("left")
    subject = _style("subject", fontName="Helvetica-Bold")
    body = _style(
        "body", TA_JUSTIFY, leading=FONT_SIZE * LINE_SPACING, spaceBefore=GAPS["paragraph"]
    )
    bullet = _style(
        "bullet", TA_JUSTIFY, leading=FONT_SIZE * LIST_LINE_SPACING, leftIndent=18, bulletIndent=6
    )
    flowables: list = [Paragraph(escape(line), right) for line in layout.sender]
    flowables.append(Spacer(0, GAPS["recipient"]))
    flowables += [Paragraph(escape(line), left) for line in layout.recipient]
    flowables += [Spacer(0, GAPS["date"]), Paragraph(escape(layout.date), right)]
    flowables += [Spacer(0, GAPS["subject"]), Paragraph(escape(layout.subject), subject)]
    flowables += [Spacer(0, GAPS["salutation"]), Paragraph(escape(layout.salutation), left)]
    for line in layout.body:
        if line.startswith(BULLET):
            flowables.append(Paragraph(escape(line.removeprefix(BULLET)), bullet, bulletText="•"))
        else:
            flowables.append(Paragraph(escape(line), body))
    flowables += [Spacer(0, GAPS["closing"]), Paragraph(escape(layout.closing), left)]
    flowables += [Spacer(0, GAPS["signature"]), Paragraph(escape(layout.signature), left)]
    return flowables
